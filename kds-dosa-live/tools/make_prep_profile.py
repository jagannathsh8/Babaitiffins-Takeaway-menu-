#!/usr/bin/env python3
"""Build assets/prep_profile.json for PREP LIVE from the BOM and Petpooja item-wise reports.

    python3 make_prep_profile.py BOM.xlsx out.json REPORT1.xlsx [REPORT2.xlsx ...]

Only quantities are kept: per-portion usage of each tracked prepared item, the raw-ingredient
recipe (per kg) of each tracked item, and average usage per weekday/hour. No prices, costs,
KOT ids, tables or staff names leave the source files.
"""
import collections
import datetime
import json
import re
import statistics
import sys

import openpyxl

# Display name -> BOM product names that count as that item
TARGETS = [
    ("Sambar", ["CP - Breakfast Sambar"]),
    ("Peanut Chutney", ["OP - Peanut chutney"]),
    ("Aloo Masala", ["OP - Aloo masala"]),
    ("Sagu (Aloo Curry)", ["OP - Aloo curry"]),
    ("Upma", ["OP - Upma"]),
    ("Pongal", ["OP - New Pongal", "OP - Pongal"]),
    ("Allam Chutney", ["CP - Allam Chutney"]),
    ("Tomato Chutney", ["OP - Tomato chutney"]),
    ("Kesari Bath", ["OP - Kesari bath"]),
    ("Dosa Batter", ["OP - Dosa batter"]),
    ("Idli Batter", ["OP - Idli batter"]),
]


def norm(s):
    return re.sub(r"\s+", " ", str(s).strip().lower())


def load_bom(path):
    bom = collections.defaultdict(list)
    uom = {}
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] and r[2]:
            bom[norm(r[1])].append((r[2].strip(), float(r[8] or 0), (r[4] or "").strip()))
            uom[norm(r[1])] = (r[12] or "").strip()
    return bom, uom


def main():
    bom_path, out_path, reports = sys.argv[1], sys.argv[2], sys.argv[3:]
    bom, uom = load_bom(bom_path)
    target_of = {norm(c): i for i, (_, cs) in enumerate(TARGETS) for c in cs}

    memo = {}

    def expand(name, depth=0):
        """Per 1 unit of `name`: usage of each target (recursing through non-target OP items)."""
        key = norm(name)
        if key in memo:
            return memo[key]
        out = collections.Counter()
        if depth < 6:
            for ing, q, _ in bom.get(key, []):
                k = norm(ing)
                if k in target_of:
                    out[target_of[k]] += q
                elif k in bom and k != key:
                    for t, v in expand(ing, depth + 1).items():
                        out[t] += q * v
        memo[key] = out
        return out

    # menu item -> {target index: kg per portion}
    menu = {}
    for key in bom:
        if key.startswith(("op - ", "cp - ")):
            continue
        e = expand(key)
        if e:
            menu[key] = {str(t): round(v, 5) for t, v in e.items() if v > 0}

    # recipe breakdown per kg of each target (first BOM product of the target)
    recipes = []
    for _, cs in TARGETS:
        rows = bom.get(norm(cs[0]), [])
        recipes.append([[ing, round(q, 5), u] for ing, q, u in rows if q > 0])

    # usage per (date, hour) from the reports
    use = collections.defaultdict(lambda: [0.0] * len(TARGETS))
    seen = set()
    dosa_prep = collections.defaultdict(list)
    for path in reports:
        ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i < 6 or not r or r[0] is None or r[8] is None:
                continue
            if len(r) > 7 and r[7] not in (None, "Success"):
                continue
            t = datetime.datetime.strptime(str(r[8])[:19], "%Y-%m-%d %H:%M:%S")
            row_key = (r[0], str(r[8]), r[4], r[5])
            if row_key in seen:  # overlapping exports
                continue
            seen.add(row_key)
            for ti, v in menu.get(norm(r[4]), {}).items():
                use[(t.date(), t.hour)][int(ti)] += v * float(r[5] or 0)
            if r[1] == "Dine In" and "dosa" in str(r[4]).lower() and len(r) > 10 and r[10] is not None:
                dosa_prep[(r[0], t.date())].append(float(r[10]))

    days = sorted({d for d, _ in use})
    full_days = days[1:-1] if len(days) > 2 else days  # first/last export day may be partial
    by_wd = collections.defaultdict(list)
    for d in full_days:
        by_wd[d.weekday()].append(d)
    # profile[target][weekday][hour] = average kg (weekday 0 = Monday)
    profile = []
    for ti in range(len(TARGETS)):
        overall = [statistics.mean(use[(d, h)][ti] for d in full_days) for h in range(24)]
        wk = []
        for wd in range(7):
            ds = by_wd.get(wd)
            wk.append([round(statistics.mean(use[(d, h)][ti] for d in ds), 3) if ds else round(overall[h], 3)
                       for h in range(24)])
        profile.append(wk)

    preps = [max(v) for v in dosa_prep.values()]
    out = {
        "targets": [n for n, _ in TARGETS],
        "units": ["kg"] * len(TARGETS),
        "menu": menu,
        "recipes": recipes,
        "profile": profile,
        "profileDays": [str(d) for d in full_days],
        "dosaDineInAvgPrepMin": round(statistics.mean(preps), 2) if preps else None,
    }
    with open(out_path, "w") as f:
        json.dump(out, f, separators=(",", ":"))
    print(f"{len(menu)} menu items, {len(full_days)} profile days ({full_days[0]}..{full_days[-1]}), "
          f"dosa avg prep {out['dosaDineInAvgPrepMin']} min -> {out_path}")


if __name__ == "__main__":
    main()
